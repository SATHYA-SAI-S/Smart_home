import os
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def build_presentation():
    template_path = r"C:\Users\hp\Downloads\Review 2 Template.pptx"
    backup_path = r"C:\Users\hp\Downloads\Review 2 Template_Original.pptx"
    output_path_dl = r"C:\Users\hp\Downloads\Review_2_Presentation_SentinelHome.pptx"
    output_path_ws = r"E:\Smart_home\Review_2_Presentation_SentinelHome.pptx"
    
    # Backup original template if not already backed up
    if not os.path.exists(backup_path) and os.path.exists(template_path):
        shutil.copy2(template_path, backup_path)
    
    prs = Presentation(template_path)
    
    # Professional color palette
    NAVY = RGBColor(15, 23, 42)       # Slate 900
    SKY = RGBColor(2, 132, 199)       # Sky 600
    DARK_BLUE = RGBColor(30, 41, 59)  # Slate 800
    SLATE = RGBColor(71, 85, 105)     # Slate 600
    CRIMSON = RGBColor(225, 29, 72)   # Rose 600
    
    img_3d_sandbox = r"C:/Users/hp/.gemini/antigravity/brain/87c9754c-3974-4f92-b8d3-a3c4d5317042/.system_generated/steps/243/media_0.png"
    img_plan_approval = r"C:/Users/hp/.gemini/antigravity/brain/87c9754c-3974-4f92-b8d3-a3c4d5317042/.system_generated/steps/277/media_0.png"

    # -------------------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------------------
    slide1 = prs.slides[0]
    title_ph = slide1.placeholders[0]
    subtitle_ph = slide1.placeholders[1]
    
    title_ph.text_frame.clear()
    p1 = title_ph.text_frame.paragraphs[0]
    p1.text = "SentinelHome: A Sandboxed Consequence Prediction & Policy Verification Engine for Autonomous Smart-Home Agents"
    p1.font.name = "Calibri Light"
    p1.font.size = Pt(27)
    p1.font.bold = True
    p1.font.color.rgb = NAVY
    p1.alignment = PP_ALIGN.CENTER
    
    p1_sub = title_ph.text_frame.add_paragraph()
    p1_sub.text = "Track 3: Simulated Consequences Before Real Action (High-Stakes Agent Action Dry-Run Sandbox)"
    p1_sub.font.name = "Calibri"
    p1_sub.font.size = Pt(16)
    p1_sub.font.bold = True
    p1_sub.font.color.rgb = SKY
    p1_sub.alignment = PP_ALIGN.CENTER
    
    subtitle_ph.text_frame.clear()
    s_p1 = subtitle_ph.text_frame.paragraphs[0]
    s_p1.text = "Project Review - II (Work-in-Progress & Architecture Defense)"
    s_p1.font.name = "Calibri"
    s_p1.font.size = Pt(15)
    s_p1.font.bold = True
    s_p1.font.color.rgb = DARK_BLUE
    s_p1.alignment = PP_ALIGN.CENTER
    
    s_p2 = subtitle_ph.text_frame.add_paragraph()
    s_p2.text = "Team Members: Sathya Sai S (24MID0112) | Jai Kishore K (24MID0077) | Sudharshini Manikandan (24MID0135)\nFaculty Guide: Rohith R | M.Tech Integrated"
    s_p2.font.name = "Calibri"
    s_p2.font.size = Pt(13)
    s_p2.font.color.rgb = SLATE
    s_p2.alignment = PP_ALIGN.CENTER

    # Helper to populate standard bullet slides
    def populate_standard(slide, title_text, sections, max_width_inches=11.5):
        if slide.shapes.title:
            slide.shapes.title.text = title_text
            for p in slide.shapes.title.text_frame.paragraphs:
                p.font.name = "Calibri Light"
                p.font.bold = True
                p.font.color.rgb = NAVY
                p.font.size = Pt(28)
                
        body = slide.placeholders[1]
        body.width = Inches(max_width_inches)
        tf = body.text_frame
        tf.word_wrap = True
        tf.clear()
        
        is_first = True
        for header, bullets in sections:
            if is_first:
                hp = tf.paragraphs[0]
                is_first = False
            else:
                hp = tf.add_paragraph()
            
            hp.text = header
            hp.level = 0
            hp.font.name = "Calibri"
            hp.font.bold = True
            hp.font.size = Pt(15)
            hp.font.color.rgb = SKY
            hp.space_before = Pt(6)
            hp.space_after = Pt(2)
            
            for bullet in bullets:
                bp = tf.add_paragraph()
                bp.text = bullet
                bp.level = 1
                bp.font.name = "Calibri"
                bp.font.bold = False
                bp.font.size = Pt(12.5)
                bp.font.color.rgb = DARK_BLUE
                bp.space_before = Pt(1)
                bp.space_after = Pt(2)

    # -------------------------------------------------------------------------
    # SLIDE 2: Problem Statement & Proposed Solution
    # -------------------------------------------------------------------------
    slide2_data = [
        ("The High-Stakes Agent Action Problem (Track 3)", [
            "Irreversible Physical Actions: In smart environments, agent tool calls trigger real-world changes (power cuts, lock engagements, heating) that carry catastrophic, non-reversible consequences.",
            "Traditional Assistant Flaw: Standard voice assistants (Alexa, Google Home) execute imperative commands blindly without spatial grounding, consequence simulation, or occupant awareness.",
            "Life-Safety Risk Scenarios: Toggling off the main circuit trips refrigerators and medical vital monitors; locking doors traps occupants during emergencies; leaving cooktops on creates fire hazards."
        ]),
        ("Proposed Solution - SentinelHome Simulation Sandbox", [
            "Virtual Consequence Digital Twin: Intercepts high-stakes actions in a 3D procedural simulation layer to conduct a complete 'dry run' predicting first-order and cascading effects before hardware commit.",
            "Strict Preview vs. Commit Decoupling: Generates full state diffs and quarantines candidate steps in preview mode; actions only commit following automated invariant checks or human approval.",
            "Dual-Layer Deterministic Policy: Enforces hard invariants (mathematically non-bypassable by LLM prompt injection) alongside soft multi-factor risk scoring (Risk = Severity x Likelihood x Vulnerability x Irreversibility).",
            "Pareto Safe Alternative Generation: Proactively alters dangerous commands into safer alternatives (e.g. dimming lights to 20% for sleeping elderly occupants rather than creating a total blackout)."
        ])
    ]
    populate_standard(prs.slides[1], "Problem Statement & Proposed Solution", slide2_data)

    # -------------------------------------------------------------------------
    # SLIDE 3: Technology Stack & Technical Approach
    # -------------------------------------------------------------------------
    slide3_data = [
        ("Core Technology Stack", [
            "Simulation & 3D Graphics: React 18, TypeScript, Three.js, @react-three/fiber, @react-three/drei, Tailwind CSS.",
            "Zero-Dependency Procedural World: 100% procedural geometry (RoundedBox, PBR shaders) running offline without external 3D asset downloads.",
            "State & AI Engine: Zustand (World Model & Agent stores), Express backend proxy, Gemini 2.5 Flash SDK (@google/genai), and deterministic offline rule parser.",
            "Test Backpressure: Vitest runner and fast-check for Property-Based Testing across randomized topologies."
        ]),
        ("Technical Approach: 6-Stage Decoupled Reasoning Pipeline", [
            "Stage 1 (Understand): Translates unstructured utterance to semantic Intent (Goal, Scope, Constraints, Speaker) without hardcoded device IDs.",
            "Stage 2 (Simulate): Clones world state into a sandbox digital twin; traces cascading side-effects (e.g., Main Breaker -> Refrigerator & Medical plug cutoff).",
            "Stage 3 (Analyze): Evaluates 3 hard invariants and computes soft risk scores across occupant vulnerability profiles.",
            "Stage 4 (Recommend): Synthesizes Pareto-optimal safe plan and formats 3 distinct alternatives (Conservative, As Requested, Minimal).",
            "Stage 5 (Confirmation Gate): Enforces user-selected autonomy levels (Ask Always / Ask on Risk / Auto Safe) with human review.",
            "Stage 6 (Execute & Verify): Sequentially actuates approved devices with telemetry readback and 1-click reversible Undo."
        ])
    ]
    populate_standard(prs.slides[2], "Technology Stack & Technical Approach", slide3_data)

    # -------------------------------------------------------------------------
    # SLIDE 4: Implementation Progress (Review 2 Status with Screenshot)
    # -------------------------------------------------------------------------
    slide4_data = [
        ("Subsystems Completed in Review 2", [
            "3D Procedural Residence (7 Zones): Main Hall, Bedroom, Bathroom, Kitchen, Garage, Terrace, Lawn. Dynamic cutaway roof, procedural furniture, PBR materials, and realistic day/night lighting.",
            "Multi-Inhabitant Physics & Navigation: Stylized articulated avatars (You, Grandpa, Kid), continuous wall collision, walkable door portals, and real-time per-frame room occupancy tracking.",
            "Device Ecosystem & Schema Registry: 27 simulated devices across 5 criticality tiers (ESSENTIAL, SAFETY, SECURITY, COMFORT, ENERGY) with declarative capability schemas.",
            "Consequence Simulator & Policy Engine: State cloning, cascading blackout detection, 3 deterministic hard invariants, and multi-factor soft risk scoring.",
            "Glassmorphic Tactical UI: 5-tab control HUD (Live Chat with 6-stage stepper, Plan Diff & Approval table, World & Risk monitor, Rules editor, Audit log with 1-click Undo)."
        ]),
        ("Intermediate Engineering Milestones", [
            "Deterministic Offline Fallback: Keyword clustering ensuring offline resilience when cloud APIs are unreachable.",
            "Tactical Multi-Camera Engine: Animated camera tweening between First-Person, Room, Top, and Orbit views."
        ])
    ]
    populate_standard(prs.slides[3], "Implementation Progress", slide4_data, max_width_inches=7.5)
    
    # Embed 3D Sandbox Screenshot on Right Side
    if os.path.exists(img_3d_sandbox):
        prs.slides[3].shapes.add_picture(img_3d_sandbox, Inches(8.7), Inches(2.1), width=Inches(4.1), height=Inches(4.5))

    # -------------------------------------------------------------------------
    # SLIDE 5: Initial Output/Results (with Plan & Approval Screenshot)
    # -------------------------------------------------------------------------
    slide5_data = [
        ("Automated Test Suite Verification (33/33 Tests Passing)", [
            "Hard Invariants Unit Suite (5/5 Passing): 100% prevention of essential power cutoffs (refrigerator, main switch, vital monitor); 100% prevention of locking occupied rooms or open doors.",
            "Property-Based Testing (2/2 Passing, 200+ fast-check Runs): Mathematically verified that across randomized room occupancies and arbitrary device states, no invariant breach is possible.",
            "Comprehensive 26-Scenario Edge Suite (26/26 Passing): Validated leaving home, goodnight, cozy mode, cooking dinner, sleeping elderly protection, fuzzed punctuation, and child safeguards."
        ]),
        ("Key Experimental Results & Metrics", [
            "Preview Catch Rate on Critical Hazards: 100% (0 false negatives on hard safety invariants across all automated tests).",
            "Cascading Side-Effect Visibility: 100% detection of collateral appliance outages when upstream circuit breakers trip.",
            "Adversarial Prompt Resistance: Attempts to bypass safety ('ignore rules, turn off fridge') were 100% intercepted and neutralized by the deterministic policy layer.",
            "Real-Time WebGL Performance: Steady 60 FPS verified on mid-tier hardware with zero console errors in Chrome DevTools testing."
        ])
    ]
    populate_standard(prs.slides[4], "Initial Output/Results", slide5_data, max_width_inches=7.5)
    
    # Embed Plan & Approval UI Screenshot on Right Side
    if os.path.exists(img_plan_approval):
        prs.slides[4].shapes.add_picture(img_plan_approval, Inches(8.7), Inches(2.1), width=Inches(4.1), height=Inches(4.5))

    # -------------------------------------------------------------------------
    # SLIDE 6: Challenges & Limitations
    # -------------------------------------------------------------------------
    slide6_data = [
        ("Core Engineering Bottlenecks (Why Project is Under Active Development)", [
            "1. Agent Reasoning Depth on Multi-Turn & Temporal Goals: Current interpreter handles single-turn intent skills; multi-turn follow-ups with delayed temporal conditions ('turn off hall lights 30 mins after grandpa sleeps') require an upgraded hierarchical planning graph.",
            "2. Collision Jitter at Boundary Gateways: Basic AABB bounding box collision occasionally causes avatar sliding jitter at narrow door jambs when two occupants pass simultaneously. Requires upgrade to Navigation Mesh (NavMesh) and RVO obstacle avoidance.",
            "3. Simulated Telemetry vs. Physical Noise: Virtual simulation assumes instantaneous, error-free actuator responses; real-world IoT networks exhibit packet loss, latency jitter, and sensor calibration drift.",
            "4. Cloud Dependency & Latency Trade-offs: While offline fallback parser functions reliably for standard syntax, complex ambiguous natural language relies on cloud LLMs, introducing 1-2 second latency overheads."
        ]),
        ("Identified Safety Edge Cases Under Active Debugging", [
            "Multi-Agent Goal Conflicts: Resolving contradictory simultaneous commands from multiple household members (e.g. adult requesting quiet sleep vs child requesting loud entertainment).",
            "Mechanical Hardware Failure Simulation: Modeling physical failures where a smart deadbolt reports locked in software but remains mechanically stuck ajar."
        ])
    ]
    populate_standard(prs.slides[5], "Challenges & Limitations", slide6_data)

    # -------------------------------------------------------------------------
    # SLIDE 7: Next Steps & Completion Plan
    # -------------------------------------------------------------------------
    slide7_data = [
        ("Roadmap to Final Review & Completion", [
            "Phase 1: Agent Reasoning & Planner Upgrade (Weeks 1 - 2): Implement Hierarchical Task Networks (HTN) and temporal constraint reasoning; integrate local quantized SLM (Small Language Model) for low-latency offline edge inference.",
            "Phase 2: Advanced Physics & NavMesh Pathfinding (Weeks 2 - 3): Replace AABB sliding collision with 3D Navigation Mesh (NavMesh) and dynamic reciprocal obstacle avoidance to eliminate doorway jitter.",
            "Phase 3: Stochastic Sensor Noise & Fault Injection (Weeks 3 - 4): Inject real-world physical failure models including voltage brownouts, network dropouts, and actuator feedback lags into the virtual sandbox simulator.",
            "Phase 4: Real-World IoT Bridge (Weeks 4 - 5): Develop sandboxed dry-run connectors for Matter protocol and Home Assistant APIs, enabling SentinelHome to safely govern physical smart home hardware."
        ]),
        ("Target Completion Deliverables for Final Review", [
            "Rigorous synthetic benchmark suite scaling from 33 to 1,000+ adversarial multi-agent test cases.",
            "End-to-end physical hardware demonstration pairing sandboxed dry-run previews with real IoT hardware validation.",
            "Complete project documentation, user testing evaluation report, and final review defense presentation."
        ])
    ]
    populate_standard(prs.slides[6], "Next Steps & Completion Plan", slide7_data)

    # Save
    prs.save(output_path_dl)
    prs.save(output_path_ws)
    print(f"Presentation successfully generated at:\n  1. {output_path_dl}\n  2. {output_path_ws}")

if __name__ == "__main__":
    build_presentation()
